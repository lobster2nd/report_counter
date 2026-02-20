import os
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Border, Side, Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from fields import values

RESEARCH_TYPES = [v[0].value for v in values]

COLUMN_CONFIG = {
    'A': {'name': 'Дата', 'width': 18, 'type': 'date'},
}

# Начинаем с индекса 1, так как колонка A уже занята
col_index = 1
RESEARCH_COL_MAP = {}

for research_name in RESEARCH_TYPES:
    scan_col = get_column_letter(col_index + 1)
    img_col = get_column_letter(col_index + 2)

    COLUMN_CONFIG[scan_col] = {
        'name': research_name,
        'width': 10,
        'type': 'scan',
        'group': research_name
    }
    COLUMN_CONFIG[img_col] = {
        'name': research_name,
        'width': 10,
        'type': 'img',
        'group': research_name
    }

    RESEARCH_COL_MAP[research_name] = (scan_col, img_col)
    col_index += 2

# Колонки для итогов
TOTAL_SCAN_COL = get_column_letter(col_index + 1)
TOTAL_IMG_COL = get_column_letter(col_index + 2)

COLUMN_CONFIG[TOTAL_SCAN_COL] = {
    'name': 'Итого',
    'width': 12,
    'type': 'total_scan',
    'group': 'Итого'
}
COLUMN_CONFIG[TOTAL_IMG_COL] = {
    'name': 'Итого',
    'width': 12,
    'type': 'total_img',
    'group': 'Итого'
}

MONTHS = [
    'Январь', 'Февраль', 'Март', 'Апрель', 'Май', 'Июнь',
    'Июль', 'Август', 'Сентябрь', 'Октябрь', 'Ноябрь', 'Декабрь'
]


def get_days_in_month(year: int, month_num: int) -> int:
    month_days = {
        1: 31, 2: 28, 3: 31, 4: 30, 5: 31, 6: 30,
        7: 31, 8: 31, 9: 30, 10: 31, 11: 30, 12: 31
    }

    if month_num == 2:
        if (year % 4 == 0 and year % 100 != 0) or (year % 400 == 0):
            return 29

    return month_days.get(month_num, 30)


def get_month_sheet_name(year: int, month: str) -> str:
    return f"{month}_{year}"


def get_year_file_path(year: int) -> str:
    return f'Отчёт_по_исследованиям_{year}_год.xlsx'


def get_column_index(col_letter: str) -> int:
    result = 0
    for char in col_letter:
        result = result * 26 + (ord(char) - ord('A') + 1)
    return result


def create_month_sheet(wb, year: int, month: str, month_num: int):
    sheet_name = get_month_sheet_name(year, month)

    if sheet_name in wb.sheetnames:
        return wb[sheet_name]

    sheet = wb.create_sheet(sheet_name)

    # Устанавливаем ширину колонок
    for col, config in COLUMN_CONFIG.items():
        sheet.column_dimensions[col].width = config['width']

    # Стили
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4',
                              fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center',
                                 wrap_text=True)

    subheader_font = Font(bold=True, size=10)
    subheader_fill = PatternFill(start_color='5B9BD5', end_color='5B9BD5',
                                 fill_type='solid')

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    thick_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thick'),
        bottom=Side(style='thin')
    )

    # Сортируем колонки для правильного отображения
    sorted_cols = sorted(COLUMN_CONFIG.keys(), key=get_column_index)

    # Сначала заполняем все ячейки значениями
    # Заполняем первую строку
    for col in sorted_cols:
        config = COLUMN_CONFIG[col]
        cell = sheet[f'{col}1']

        if config['type'] == 'date':
            cell.value = 'Дата'
        elif config['type'] in ['scan', 'img']:
            cell.value = config['group']
        elif config['type'] in ['total_scan', 'total_img']:
            if config['type'] == 'total_scan':
                cell.value = 'Итого за день'

        if cell.value:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

    # Заполняем вторую строку
    for col, config in COLUMN_CONFIG.items():
        cell = sheet[f'{col}2']

        if config['type'] == 'date':
            cell.value = ''
        elif config['type'] == 'scan':
            cell.value = 'иссл.'
        elif config['type'] == 'img':
            cell.value = 'снимки'
        elif config['type'] == 'total_scan':
            cell.value = 'иссл.'
        elif config['type'] == 'total_img':
            cell.value = 'снимки'

        if config['type'] != 'date':
            cell.font = subheader_font
            cell.fill = subheader_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        else:
            cell.border = thin_border

    # Объединяем ячейки
    sheet.merge_cells('A1:A2')

    groups = {}
    for col in sorted_cols:
        config = COLUMN_CONFIG[col]
        if config['type'] in ['scan', 'img']:
            if config['group'] not in groups:
                groups[config['group']] = []
            groups[config['group']].append(col)

    for group_name, cols in groups.items():
        if cols and len(cols) > 1:
            start_col = cols[0]
            end_col = cols[-1]
            sheet.merge_cells(f'{start_col}1:{end_col}1')

    sheet.merge_cells(f'{TOTAL_SCAN_COL}1:{TOTAL_IMG_COL}1')

    # Заполняем дни месяца
    days_in_month = get_days_in_month(year, month_num)
    scan_cols = [col for col, cfg in COLUMN_CONFIG.items() if
                 cfg.get('type') == 'scan']
    img_cols = [col for col, cfg in COLUMN_CONFIG.items() if
                cfg.get('type') == 'img']

    for day in range(1, days_in_month + 1):
        row = day + 2

        date_cell = sheet[f'A{row}']
        date_cell.value = f'{day:02d}.{month_num:02d}.{year}'
        date_cell.border = thin_border
        date_cell.alignment = Alignment(horizontal='center', vertical='center')

        for col in COLUMN_CONFIG.keys():
            cell = sheet[f'{col}{row}']
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')

        if scan_cols:
            scan_formula = f'=SUM({",".join([f"{col}{row}" for col in scan_cols])})'
            sheet[f'{TOTAL_SCAN_COL}{row}'].value = scan_formula
        if img_cols:
            img_formula = f'=SUM({",".join([f"{col}{row}" for col in img_cols])})'
            sheet[f'{TOTAL_IMG_COL}{row}'].value = img_formula

    # Добавляем строку с итогами за месяц
    total_row = days_in_month + 4

    for col in COLUMN_CONFIG.keys():
        if col != 'A':
            sheet[f'{col}{total_row - 1}'].border = thick_border

    sheet[f'A{total_row}'].value = 'Итого за месяц:'
    sheet[f'A{total_row}'].font = Font(bold=True)
    sheet[f'A{total_row}'].alignment = Alignment(horizontal='right')
    sheet[f'A{total_row}'].border = thick_border
    sheet.row_dimensions[total_row].height = 20

    for col in COLUMN_CONFIG.keys():
        if col in [TOTAL_SCAN_COL, TOTAL_IMG_COL]:
            formula = f'=SUM({col}3:{col}{days_in_month + 2})'
            sheet[f'{col}{total_row}'].value = formula
            sheet[f'{col}{total_row}'].font = Font(bold=True)
            sheet[f'{col}{total_row}'].fill = PatternFill(start_color='FFC000',
                                                          end_color='FFC000',
                                                          fill_type='solid')
        elif col != 'A':
            formula = f'=SUM({col}3:{col}{days_in_month + 2})'
            sheet[f'{col}{total_row}'].value = formula
            sheet[f'{col}{total_row}'].font = Font(bold=True)

        cell = sheet[f'{col}{total_row}']
        cell.border = thick_border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    sheet.freeze_panes = 'A3'
    return sheet


def create_yearly_summary(wb, year: int):
    """Создает годовой отчет на отдельном листе с диаграммой"""
    summary_sheet_name = f"Годовой_отчёт_{year}"

    # Если лист уже существует, удаляем его и создаем заново
    if summary_sheet_name in wb.sheetnames:
        wb.remove(wb[summary_sheet_name])

    # Создаем новый лист в начале книги
    wb.create_sheet(summary_sheet_name, 0)
    sheet = wb[summary_sheet_name]

    # Заголовок
    sheet.merge_cells('A1:C1')
    title_cell = sheet['A1']
    title_cell.value = f'Отчёт по исследованиям {year} год'
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal='center')

    # Стили
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4',
                              fill_type='solid')
    subheader_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2',
                                 fill_type='solid')
    bold_font = Font(bold=True)

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Устанавливаем ширину колонок
    sheet.column_dimensions['A'].width = 40
    sheet.column_dimensions['B'].width = 20
    sheet.column_dimensions['C'].width = 20

    # Заголовки таблицы
    headers = ['Наименование', 'Всего исследований', 'Всего снимков']
    for i, header in enumerate(headers, start=1):
        cell = sheet.cell(row=3, column=i)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    # Собираем данные по месяцам
    month_totals = {}
    monthly_data = []

    # Инициализируем totals для всех типов исследований из RESEARCH_TYPES
    for research_name in RESEARCH_TYPES:
        month_totals[research_name] = {'scan': 0, 'img': 0}

    # Добавляем группы
    month_totals['Костно-мышечной системы'] = {'scan': 0, 'img': 0}
    month_totals['Черепа и челюстно-лицевой области'] = {'scan': 0, 'img': 0}

    # Проходим по всем месяцам для сбора данных
    for month_num, month_name in enumerate(MONTHS, start=1):
        month_sheet_name = get_month_sheet_name(year, month_name)
        month_total = 0

        if month_sheet_name in wb.sheetnames:
            month_sheet = wb[month_sheet_name]
            days_in_month = get_days_in_month(year, month_num)

            # Суммируем по дням месяца
            for day in range(1, days_in_month + 1):
                row = day + 2

                for research_name in RESEARCH_TYPES:
                    scan_col, img_col = RESEARCH_COL_MAP.get(research_name,
                                                             (None, None))
                    if scan_col and img_col:
                        scan_cell = month_sheet[f'{scan_col}{row}']

                        if scan_cell.value and isinstance(scan_cell.value,
                                                          (int, float)):
                            month_totals[research_name][
                                'scan'] += scan_cell.value
                            month_total += scan_cell.value

                        if img_cell.value and isinstance(img_cell.value, (int, float)):
                            month_totals[research_name]['img'] += img_cell.value

        monthly_data.append(month_total)

    # Вычисляем суммы для групп
    # Костно-мышечной системы
    musculoskeletal_groups = [
        'Конечности',
        'Таза и тазобедренных суставов',
        'Шейные позвонки',
        'Грудные позвонки',
        'Поясничные позвонки',
        'Рёбра и грудина'
    ]

    for group in musculoskeletal_groups:
        if group in month_totals:
            month_totals['Костно-мышечной системы']['scan'] += \
            month_totals[group]['scan']

    # Черепа и челюстно-лицевой области
    skull_groups = [
        'Зубы',
        'Челюстей',
        'Околоносовых пазух',
        'Череп'
    ]

    for group in skull_groups:
        if group in month_totals:
            month_totals['Черепа и челюстно-лицевой области']['scan'] += \
            month_totals[group]['scan']

    # Заполняем данные по категориям
    row = 4
    total_scan_all = 0
    total_img_all = 0

    # Считаем общий итог (только по основным категориям, без групп)
    for research_name in RESEARCH_TYPES:
        scan_total = month_totals.get(research_name, {}).get('scan', 0)
        img_total = month_totals.get(research_name, {}).get('img', 0)
        total_scan_all += scan_total
        total_img_all += img_total

    # Строка "Всего" (жирным шрифтом)
    sheet.cell(row=row, column=1, value='Всего')
    sheet.cell(row=row, column=2, value=total_scan_all)
    sheet.cell(row=row, column=3, value=total_img_all)

    for col in range(1, 4):
        cell = sheet.cell(row=row, column=col)
        cell.font = Font(bold=True)
        cell.fill = subheader_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')

    row += 1

    # Точный порядок строк
    ordered_categories = [
        'ОГК',
        'Костно-мышечной системы',  # Группа (жирным)
        'Конечности',
        'Таза и тазобедренных суставов',
        'Шейные позвонки',
        'Грудные позвонки',
        'Поясничные позвонки',
        'Рёбра и грудина',
        'Черепа и челюстно-лицевой области',  # Группа (жирным)
        'Зубы',
        'Челюстей',
        'Околоносовых пазух',
        'Череп',
        'Брюшная полость'
    ]

    # Заполняем данные в заданном порядке
    for category in ordered_categories:
        if category in month_totals:
            scan_total = month_totals[category]['scan']
            img_total = month_totals[category]['img']

            sheet.cell(row=row, column=1, value=category)
            sheet.cell(row=row, column=2, value=scan_total)
            sheet.cell(row=row, column=3, value=img_total)

            for col in range(1, 4):
                cell = sheet.cell(row=row, column=col)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center')

                # Делаем жирным шрифт для групп
                if category in ['Костно-мышечной системы',
                                'Черепа и челюстно-лицевой области']:
                    cell.font = bold_font

            row += 1

    # Создаем данные для диаграммы на отдельном листе
    chart_data_sheet = wb.create_sheet("_chart_data", 1)

    # Заголовки
    chart_data_sheet['A1'] = 'Месяц'
    chart_data_sheet['B1'] = 'Количество исследований'

    # Заполняем данные по месяцам
    for i, (month_name, month_value) in enumerate(zip(MONTHS, monthly_data),
                                                  start=2):
        chart_data_sheet[f'A{i}'] = month_name
        chart_data_sheet[f'B{i}'] = month_value

    # Создаем столбчатую диаграмму
    from openpyxl.chart import BarChart, Reference

    chart = BarChart()
    chart.title = f"Количество исследований по месяцам"
    chart.style = 12
    chart.y_axis.title = "Количество исследований"
    chart.height = 9
    chart.width = 20
    chart.shape = 4  # Прямоугольные столбцы
    chart.legend = None

    # Данные для диаграммы
    data = Reference(chart_data_sheet, min_col=2, min_row=1, max_row=13)
    categories = Reference(chart_data_sheet, min_col=1, min_row=2, max_row=13)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    # Добавляем диаграмму на основной лист
    sheet.add_chart(chart, "E2")

    # Скрываем лист с данными для диаграммы
    chart_data_sheet.sheet_state = 'hidden'

    # Добавляем дату формирования отчета
    sheet.merge_cells(f'A{row + 4}:C{row + 4}')
    date_cell = sheet.cell(row=row + 4, column=1)
    date_cell.value = f'Сформировано: {datetime.now().strftime("%d.%m.%Y %H:%M")}'
    date_cell.font = Font(italic=True)
    date_cell.alignment = Alignment(horizontal='right')

    return sheet

def create_table(year: int = None):
    if year is None:
        year = datetime.now().year

    file_path = get_year_file_path(year)

    if os.path.exists(file_path):
        wb = openpyxl.load_workbook(file_path)
    else:
        wb = openpyxl.Workbook()
        if 'Sheet' in wb.sheetnames:
            del wb['Sheet']

    # Создаем месячные листы
    for month_num, month_name in enumerate(MONTHS, start=1):
        create_month_sheet(wb, year, month_name, month_num)

    # Создаем годовой отчет
    create_yearly_summary(wb, year)

    wb.save(file_path)
    return file_path


def update_yearly_summary(year: int = None):
    """Обновляет годовой отчет после добавления новых данных"""
    if year is None:
        year = datetime.now().year

    file_path = get_year_file_path(year)
    if not os.path.exists(file_path):
        create_table(year)
        return

    wb = openpyxl.load_workbook(file_path)
    create_yearly_summary(wb, year)
    wb.save(file_path)
