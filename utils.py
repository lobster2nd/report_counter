from datetime import datetime
import flet as ft
import openpyxl
from fields import values
from tables import (
    MONTHS,
    get_year_file_path,
    get_month_sheet_name,
    create_month_sheet,
    RESEARCH_COL_MAP,
    update_yearly_summary
)


def show_info(info, page):
    page.snack_bar = ft.SnackBar(ft.Text(f'{info}'))
    page.snack_bar.open = True
    page.update()


def clear_fields(e, page):
    for v in values:
        v[1].value = ''
        v[2].value = ''
    page.update()


def get_month_year_from_date(date_str):
    try:
        if not date_str:
            return None, None
        date_obj = datetime.strptime(date_str, '%d.%m.%Y')
        month_num = date_obj.month
        year = date_obj.year
        return month_num, year
    except (ValueError, IndexError) as e:
        print(f"Ошибка преобразования даты: {e}")
        return None, None


def validate_data(date_str, page):
    if not date_str:
        show_info('Выберите дату', page)
        return False

    month_num, year = get_month_year_from_date(date_str)
    if not month_num:
        show_info('Некорректный формат даты', page)
        return False

    for v in values:
        scan = v[1].value
        img = v[2].value

        if scan and img:
            try:
                int(scan)
                int(img)
            except ValueError:
                show_info('Значение должно быть целым числом', page)
                return False

        if (scan and not img) or (img and not scan):
            show_info('Введите количество исследований и снимков', page)
            return False

    return True


def find_row_by_date(sheet, date_str):
    for row in range(3, 35):
        cell_value = sheet[f'A{row}'].value
        if cell_value and date_str == str(cell_value):
            return row
    return None


def save_to_journal(date_str, page):
    month_num, year = get_month_year_from_date(date_str)
    if not month_num:
        show_info('Ошибка определения месяца', page)
        return

    month_name = MONTHS[month_num - 1]
    file_path = get_year_file_path(year)
    sheet_name = get_month_sheet_name(year, month_name)

    try:
        wb = openpyxl.load_workbook(file_path)
    except FileNotFoundError:
        from tables import create_table
        create_table(year)
        wb = openpyxl.load_workbook(file_path)

    if sheet_name not in wb.sheetnames:
        from tables import create_month_sheet
        create_month_sheet(wb, year, month_name, month_num)

    sheet = wb[sheet_name]
    row = find_row_by_date(sheet, date_str)

    if not row:
        show_info(f'Дата {date_str} не найдена в журнале', page)
        return

    # Сохраняем значения
    data_saved = False
    for val in values:
        section_name = val[0].value
        scan_cnt_str = val[1].value
        img_cnt_str = val[2].value

        if not scan_cnt_str or not img_cnt_str:
            continue

        try:
            scan_cnt = int(scan_cnt_str)
            img_cnt = int(img_cnt_str)
        except ValueError:
            continue

        col_pair = RESEARCH_COL_MAP.get(section_name)
        if col_pair:
            scan_col, img_col = col_pair
            sheet[f'{scan_col}{row}'].value = scan_cnt
            sheet[f'{img_col}{row}'].value = img_cnt
            print(
                f"✅ Сохранено: {section_name} - иссл:{scan_cnt}, снимки:{img_cnt}")  # Отладка
            data_saved = True

    if data_saved:
        # Сохраняем файл
        wb.save(file_path)
        print(f"✅ Файл сохранен: {file_path}")

        # Перезагружаем файл и обновляем годовой отчет
        wb = openpyxl.load_workbook(file_path)
        from tables import create_yearly_summary
        create_yearly_summary(wb, year)
        wb.save(file_path)
        print(f"✅ Годовой отчет обновлен")

        clear_fields(None, page)
        show_info(f'Данные за {date_str} сохранены', page)
    else:
        show_info('Нет данных для сохранения', page)


def add_to_table_values(e, page, date_str):
    if validate_data(date_str=date_str, page=page):
        save_to_journal(date_str, page)
